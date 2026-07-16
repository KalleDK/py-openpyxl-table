"""Tests for openpyxl_table core functionality."""

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from openpyxl_table import CellCord, DictWriter, _DictWriter, get_area, get_coord


# ---------------------------------------------------------------------------
# CellCord
# ---------------------------------------------------------------------------


class TestCellCord:
    def test_basic_construction(self):
        c = CellCord(column=0, row=0)
        assert c.column == 0
        assert c.row == 0

    def test_named_tuple_unpacking(self):
        col, row = CellCord(column=3, row=7)
        assert col == 3
        assert row == 7

    def test_offset_both(self):
        c = CellCord(column=1, row=2)
        assert c.offset(column=3, row=4) == CellCord(column=4, row=6)

    def test_offset_column_only(self):
        c = CellCord(column=0, row=5)
        assert c.offset(column=2) == CellCord(column=2, row=5)

    def test_offset_row_only(self):
        c = CellCord(column=3, row=0)
        assert c.offset(row=1) == CellCord(column=3, row=1)

    def test_offset_zero_is_identity(self):
        c = CellCord(column=4, row=7)
        assert c.offset() == c

    def test_from_string_a1(self):
        c = CellCord.from_string("A1")
        assert c == CellCord(column=0, row=0)

    def test_from_string_b3(self):
        c = CellCord.from_string("B3")
        assert c == CellCord(column=1, row=2)

    def test_from_string_z26(self):
        c = CellCord.from_string("Z26")
        assert c == CellCord(column=25, row=25)

    def test_from_string_aa1(self):
        # Multi-letter column
        c = CellCord.from_string("AA1")
        assert c.column == 26
        assert c.row == 0

    def test_roundtrip_from_string_to_get_coord(self):
        for addr in ("A1", "B3", "Z10", "AA5", "C100"):
            c = CellCord.from_string(addr)
            assert get_coord(c) == addr


# ---------------------------------------------------------------------------
# get_coord
# ---------------------------------------------------------------------------


class TestGetCoord:
    def test_origin(self):
        assert get_coord(CellCord(0, 0)) == "A1"

    def test_second_column(self):
        assert get_coord(CellCord(1, 0)) == "B1"

    def test_second_row(self):
        assert get_coord(CellCord(0, 1)) == "A2"

    def test_arbitrary(self):
        assert get_coord(CellCord(2, 4)) == "C5"


# ---------------------------------------------------------------------------
# get_area
# ---------------------------------------------------------------------------


class TestGetArea:
    def test_single_cell(self):
        assert get_area(CellCord(0, 0), 1, 1) == "A1:A1"

    def test_single_row(self):
        assert get_area(CellCord(0, 0), 3, 1) == "A1:C1"

    def test_single_column(self):
        assert get_area(CellCord(0, 0), 1, 3) == "A1:A3"

    def test_rectangle(self):
        assert get_area(CellCord(0, 0), 3, 2) == "A1:C2"

    def test_offset_start(self):
        assert get_area(CellCord(1, 1), 2, 2) == "B2:C3"


# ---------------------------------------------------------------------------
# _DictWriter (unit tests without openpyxl I/O)
# ---------------------------------------------------------------------------


def make_ws() -> Worksheet:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    return ws


class TestDictWriterInit:
    def test_no_auto_header(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=False)
        assert writer._height == 0

    def test_auto_header_increments_height(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=True)
        assert writer._height == 1

    def test_auto_header_writes_values(self):
        ws = make_ws()
        _DictWriter(["Name", "Score"], CellCord(0, 0), ws, auto_header=True)
        assert ws.cell(1, 1).value == "Name"
        assert ws.cell(1, 2).value == "Score"


class TestDictWriterArea:
    def test_area_no_rows(self):
        """area must return a valid (non-zero-height) reference even when nothing is written."""
        ws = make_ws()
        writer = _DictWriter(["A", "B", "C"], CellCord(0, 0), ws, auto_header=False)
        area = writer.area
        # Must not contain row "0" which would be invalid
        assert "0" not in area
        assert area == "A1:C1"

    def test_area_after_header(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=True)
        assert writer.area == "A1:B1"

    def test_area_after_header_and_one_row(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=True)
        writer.writerow({"A": 1, "B": 2})
        assert writer.area == "A1:B2"

    def test_area_offset_start(self):
        ws = make_ws()
        writer = _DictWriter(["X", "Y"], CellCord(1, 2), ws, auto_header=True)
        writer.writerow({"X": 10, "Y": 20})
        # Starts at B3 (0-based col=1, row=2 → 1-based B3), 2 rows → B3:C4
        assert writer.area == "B3:C4"


class TestDictWriterWriteRow:
    def test_basic_row(self):
        ws = make_ws()
        writer = _DictWriter(["Name", "Score"], CellCord(0, 0), ws, auto_header=False)
        writer.writerow({"Name": "Alice", "Score": 95})
        assert ws.cell(1, 1).value == "Alice"
        assert ws.cell(1, 2).value == 95
        assert writer._height == 1

    def test_missing_key_leaves_cell_empty(self):
        ws = make_ws()
        writer = _DictWriter(["Name", "Score"], CellCord(0, 0), ws, auto_header=False)
        writer.writerow({"Name": "Alice"})  # "Score" is absent
        assert ws.cell(1, 1).value == "Alice"
        assert ws.cell(1, 2).value is None

    def test_extra_key_is_ignored(self):
        ws = make_ws()
        writer = _DictWriter(["Name"], CellCord(0, 0), ws, auto_header=False)
        writer.writerow({"Name": "Alice", "Extra": "ignored"})
        assert ws.cell(1, 1).value == "Alice"

    def test_multiple_rows_stacked(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=False)
        writer.writerow({"A": 1, "B": 2})
        writer.writerow({"A": 3, "B": 4})
        assert ws.cell(1, 1).value == 1
        assert ws.cell(2, 1).value == 3
        assert writer._height == 2

    def test_tuple_fieldnames_resolved(self):
        ws = make_ws()
        writer = _DictWriter([("Score", 10)], CellCord(0, 0), ws, auto_header=False)
        writer.writerow({"Score": 42})
        assert ws.cell(1, 1).value == 42


class TestDictWriterWriteRows:
    def test_writerows_all_items(self):
        ws = make_ws()
        writer = _DictWriter(["A", "B"], CellCord(0, 0), ws, auto_header=False)
        writer.writerows([{"A": 1, "B": 2}, {"A": 3, "B": 4}])
        assert writer._height == 2

    def test_writerows_empty_list(self):
        ws = make_ws()
        writer = _DictWriter(["A"], CellCord(0, 0), ws, auto_header=False)
        writer.writerows([])
        assert writer._height == 0


class TestDictWriterWriteHeader:
    def test_header_sets_column_width(self):
        ws = make_ws()
        writer = _DictWriter([("Name", 25), ("Score", 10)], CellCord(0, 0), ws, auto_header=False)
        writer.writeheader()
        assert ws.column_dimensions["A"].width == 25
        assert ws.column_dimensions["B"].width == 10

    def test_mixed_fieldnames(self):
        ws = make_ws()
        _DictWriter(["ID", ("Name", 20), "Score"], CellCord(0, 0), ws, auto_header=True)
        assert ws.cell(1, 1).value == "ID"
        assert ws.cell(1, 2).value == "Name"
        assert ws.cell(1, 3).value == "Score"
        assert ws.column_dimensions["B"].width == 20
        # "ID" and "Score" should not set a custom column width
        assert ws.column_dimensions["A"].width != 20


# ---------------------------------------------------------------------------
# DictWriter context manager (integration)
# ---------------------------------------------------------------------------


class TestDictWriterContextManager:
    def test_table_is_registered(self):
        ws = make_ws()
        with DictWriter(ws, "A1", ["Name", "Score"], auto_header=True) as writer:
            writer.writerow({"Name": "Alice", "Score": 95})
        assert len(ws.tables) == 1

    def test_table_display_name(self):
        ws = make_ws()
        with DictWriter(ws, "A1", ["X"], displayName="MyTable") as writer:
            writer.writerow({"X": 1})
        table = list(ws.tables.values())[0]
        assert table.displayName == "MyTable"

    def test_table_ref_matches_area(self):
        ws = make_ws()
        with DictWriter(ws, "A1", ["A", "B"], auto_header=True) as writer:
            writer.writerow({"A": 1, "B": 2})
        table = list(ws.tables.values())[0]
        assert table.ref == "A1:B2"

    def test_string_cell_address_accepted(self):
        ws = make_ws()
        with DictWriter(ws, "C3", ["X", "Y"], auto_header=True) as writer:
            writer.writerow({"X": 10, "Y": 20})
        table = list(ws.tables.values())[0]
        assert table.ref == "C3:D4"

    def test_cellcord_accepted(self):
        ws = make_ws()
        cell = CellCord.from_string("B2")
        with DictWriter(ws, cell, ["X"], auto_header=True) as writer:
            writer.writerow({"X": 99})
        table = list(ws.tables.values())[0]
        assert table.ref == "B2:B3"

    def test_custom_style_applied(self):
        ws = make_ws()
        with DictWriter(ws, "A1", ["X"], style="TableStyleLight1") as writer:
            writer.writerow({"X": 1})
        table = list(ws.tables.values())[0]
        assert table.tableStyleInfo is not None
        assert table.tableStyleInfo.name == "TableStyleLight1"

    def test_no_rows_produces_valid_ref(self):
        """A table with no written rows should still register with a valid range."""
        ws = make_ws()
        with DictWriter(ws, "A1", ["Name", "Score"]):
            pass  # write nothing
        table = list(ws.tables.values())[0]
        # Table reference must be valid (must not contain row 0)
        assert "0" not in table.ref

    def test_workbook_is_saveable(self, tmp_path):
        """The produced workbook must be writable without errors."""
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        with DictWriter(ws, "A1", ["Name", "Value"], auto_header=True) as writer:
            writer.writerow({"Name": "test", "Value": 42})
        out = tmp_path / "output.xlsx"
        wb.save(str(out))
        assert out.exists()

    def test_multiple_tables_on_sheet(self):
        ws = make_ws()
        with DictWriter(ws, "A1", ["X"], auto_header=True, displayName="T1") as w:
            w.writerow({"X": 1})
        with DictWriter(ws, "C1", ["Y"], auto_header=True, displayName="T2") as w:
            w.writerow({"Y": 2})
        assert len(ws.tables) == 2
