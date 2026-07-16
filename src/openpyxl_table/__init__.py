"""openpyxl_table – a thin convenience wrapper around :class:`openpyxl.worksheet.table.Table`.

The main entry point is the :func:`DictWriter` context manager which lets you write
rows of data from plain dictionaries and automatically registers the filled range as a
formatted Excel table when the ``with`` block exits.

Typical usage::

    import openpyxl
    from openpyxl_table import DictWriter

    wb = openpyxl.Workbook()
    ws = wb.active

    with DictWriter(ws, "A1", ["Name", "Score"], auto_header=True) as writer:
        writer.writerow({"Name": "Alice", "Score": 95})
        writer.writerow({"Name": "Bob", "Score": 87})

    wb.save("output.xlsx")
"""

import contextlib
import dataclasses
from typing import Any, Iterator, NamedTuple

from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

__version__ = "1.2.7"


class CellCord(NamedTuple):
    """Zero-based (column, row) coordinate of a single worksheet cell.

    Both *column* and *row* are **0-based** integers, matching Python's
    typical indexing convention.  openpyxl uses 1-based indices internally;
    the conversion is handled by :func:`get_coord`.

    Example::

        c = CellCord(column=0, row=0)  # → cell "A1"
        c.offset(column=2, row=1)      # → CellCord(column=2, row=1)  → "C2"
    """

    column: int
    row: int

    def offset(self, column: int = 0, row: int = 0) -> "CellCord":
        """Return a new :class:`CellCord` shifted by *column* and *row* steps."""
        return CellCord(column=self.column + column, row=self.row + row)

    @classmethod
    def from_string(cls, cell: str) -> "CellCord":
        """Create a :class:`CellCord` from an Excel cell address such as ``"B3"``."""
        c, r = coordinate_from_string(cell)
        return cls(column=column_index_from_string(c) - 1, row=r - 1)


def get_coord(coord: CellCord) -> str:
    """Convert a zero-based :class:`CellCord` to an Excel cell address string.

    Example::

        get_coord(CellCord(0, 0))  # → "A1"
        get_coord(CellCord(2, 4))  # → "C5"
    """
    return f"{get_column_letter(coord.column + 1)}{coord.row + 1}"


def get_area(coord: CellCord, width: int, height: int) -> str:
    """Return the Excel range string for a rectangle starting at *coord*.

    *width* and *height* are expressed in number of columns/rows (1-based
    sizes), so a single cell has ``width=1, height=1``.

    Example::

        get_area(CellCord(0, 0), 3, 2)  # → "A1:C2"
    """
    return f"{get_coord(coord)}:{get_coord(coord.offset(column=width - 1, row=height - 1))}"


# A fieldname entry is either a plain column name string, or a (name, width) tuple
# that additionally sets the column width in the worksheet.
type Fieldnames = list[str | tuple[str, int]] | list[str] | list[tuple[str, int]]


@dataclasses.dataclass
class _DictWriter:
    """Internal writer that tracks state while rows are being appended.

    Instances are created and yielded by the :func:`DictWriter` context manager;
    callers should not instantiate this class directly.
    """

    fieldnames: Fieldnames
    """Ordered list of column names (and optional widths) that define the table schema."""

    cell: CellCord
    """Zero-based coordinate of the top-left cell of the table."""

    ws: Worksheet
    """Target worksheet."""

    auto_header: bool
    """Whether a header row was (or will be) written automatically."""

    _height: int = 0
    """Number of rows written so far (including the header row if present)."""

    def __post_init__(self) -> None:
        if self.auto_header:
            self.writeheader()

    @property
    def current_first_cell(self) -> CellCord:
        """The cell where the *next* row will be written."""
        return self.cell.offset(row=self._height)

    @property
    def area(self) -> str:
        """The Excel range reference that covers all rows written so far.

        Returns a single-row reference when nothing has been written yet so
        that an empty table still has a valid (degenerate) range.
        """
        # Ensure height is at least 1 so the area is never zero-height, which
        # would produce an invalid cell reference (e.g. "A0").
        effective_height = max(self._height, 1)
        return get_area(self.cell, len(self.fieldnames), effective_height)

    def _write_cell(self, coord: CellCord, value: Any) -> None:
        """Write *value* to the worksheet at the given zero-based *coord*."""
        self.ws.cell(row=coord.row + 1, column=coord.column + 1, value=value)

    def writeheader(self) -> None:
        """Write a header row using the configured *fieldnames*.

        Column names given as ``(name, width)`` tuples also set the
        corresponding worksheet column width.
        """
        cell = self.current_first_cell
        for i, fieldname in enumerate(self.fieldnames):
            if isinstance(fieldname, tuple):
                fieldname, width = fieldname
                self.ws.column_dimensions[get_column_letter(cell.column + i + 1)].width = width
            self._write_cell(cell.offset(column=i), fieldname)
        self._height += 1

    def writerow(self, row: dict[str, Any]) -> None:
        """Write a single data row from a dictionary keyed by column name.

        Keys that are present in *row* but not in *fieldnames* are ignored.
        Fieldnames that are absent from *row* leave the corresponding cell
        empty (no value is written).
        """
        cell = self.current_first_cell
        for i, fieldname in enumerate(self.fieldnames):
            if isinstance(fieldname, tuple):
                fieldname, _ = fieldname
            try:
                value = row[fieldname]
                self._write_cell(cell.offset(column=i), value)
            except KeyError:
                pass
        self._height += 1

    def writerows(self, rows: list[dict[str, Any]]) -> None:
        """Write multiple data rows by calling :meth:`writerow` for each item."""
        for row in rows:
            self.writerow(row)


@contextlib.contextmanager
def DictWriter(
    ws: Worksheet,
    cell: str | CellCord,
    fieldnames: Fieldnames,
    auto_header: bool = False,
    displayName: str = "Table1",
    style: str = "TableStyleMedium9",
) -> Iterator[_DictWriter]:
    """Context manager that writes dictionary rows into an Excel table.

    On entry it yields a :class:`_DictWriter` instance.  On exit it registers
    the populated range as a named, styled :class:`openpyxl.worksheet.table.Table`.

    Args:
        ws: The target :class:`~openpyxl.worksheet.worksheet.Worksheet`.
        cell: Top-left cell of the table, either as an Excel address string
            (e.g. ``"B2"``) or a :class:`CellCord`.
        fieldnames: Ordered column definitions.  Each entry is either a plain
            ``str`` column name, or a ``(name, width)`` tuple that also sets the
            worksheet column width.
        auto_header: When ``True``, a header row is written automatically as
            the first row of the table.  Defaults to ``False``.
        displayName: The name of the Excel table object.  Must be unique within
            the workbook and must not contain spaces.  Defaults to
            ``"Table1"``.
        style: The :class:`~openpyxl.worksheet.table.TableStyleInfo` name to
            apply.  Defaults to ``"TableStyleMedium9"``.

    Yields:
        A :class:`_DictWriter` that accepts :meth:`~_DictWriter.writerow` and
        :meth:`~_DictWriter.writerows` calls.

    Example::

        with DictWriter(ws, "A1", ["Name", "Age"], auto_header=True) as w:
            w.writerow({"Name": "Alice", "Age": 30})
    """
    if isinstance(cell, str):
        cell = CellCord.from_string(cell)

    writer = _DictWriter(fieldnames, cell, ws, auto_header)

    yield writer

    tab = Table(displayName=displayName, ref=writer.area)
    tab.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tab)
